from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

ZERO = Decimal("0.00")
MONEY = Decimal("0.01")

CHUNK_SIZE = 200


def _money(value: Any) -> Decimal:
    if value in (None, ""):
        return ZERO
    try:
        return Decimal(str(value)).quantize(MONEY, rounding=ROUND_HALF_UP)
    except Exception:
        return ZERO


def _amount_out(value: Decimal) -> float:
    return float(value.quantize(MONEY, rounding=ROUND_HALF_UP))


def _fetch_rows(query) -> list[dict]:
    result = query.execute()
    return list(result.data or [])


def _validate_dates(date_from: str, date_to: str) -> None:
    try:
        start = date.fromisoformat(date_from)
        end = date.fromisoformat(date_to)
    except ValueError as exc:
        raise ValueError("Dates must use YYYY-MM-DD format") from exc
    if start > end:
        raise ValueError("From date must be on or before To date")


def _chunked(items: list, size: int):
    for i in range(0, len(items), size):
        yield items[i : i + size]


SOURCE_TYPE_LABELS = {
    "bank_transaction": "Bank",
    "invoice": "Invoice",
}


def _source_label(source_type: str | None) -> str:
    if not source_type:
        return "Manual"
    return SOURCE_TYPE_LABELS.get(source_type, source_type.replace("_", " ").title())


def generate_transaction_report(
    db,
    *,
    organisation_id: str,
    date_from: str,
    date_to: str,
) -> dict:
    _validate_dates(date_from, date_to)

    # 1. Fetch posted journals in date range
    journals_raw = _fetch_rows(
        db.table("gl_journals")
        .select("id, journal_date, description, source_type, source_id, total_debit, total_credit, created_at")
        .eq("organisation_id", organisation_id)
        .eq("status", "posted")
        .gte("journal_date", date_from)
        .lte("journal_date", date_to)
        .order("journal_date")
        .order("created_at")
    )

    if not journals_raw:
        return {
            "organisation_id": organisation_id,
            "date_from": date_from,
            "date_to": date_to,
            "rows": [],
            "summary": {
                "journal_count": 0,
                "line_count": 0,
                "total_debits": 0.0,
                "total_credits": 0.0,
            },
        }

    journal_ids = [j["id"] for j in journals_raw]
    journal_map = {j["id"]: j for j in journals_raw}

    # 2. Fetch all lines for those journals (chunked to avoid URL length limits)
    lines_raw: list[dict] = []
    for chunk in _chunked(journal_ids, CHUNK_SIZE):
        lines_raw.extend(
            _fetch_rows(
                db.table("gl_journal_lines")
                .select("id, gl_journal_id, account_id, description, debit_amount, credit_amount, sort_order")
                .eq("organisation_id", organisation_id)
                .in_("gl_journal_id", chunk)
                .order("sort_order")
            )
        )

    # 3. Fetch accounts for those lines
    account_ids = list({line["account_id"] for line in lines_raw if line.get("account_id")})
    account_map: dict[str, dict] = {}
    for chunk in _chunked(account_ids, CHUNK_SIZE):
        rows = _fetch_rows(
            db.table("accounts")
            .select("id, code, name, type")
            .eq("organisation_id", organisation_id)
            .in_("id", chunk)
        )
        for row in rows:
            account_map[row["id"]] = row

    # 4. Build output rows preserving journal/sort order
    rows: list[dict] = []
    total_debits = ZERO
    total_credits = ZERO

    for journal_id in journal_ids:
        journal = journal_map[journal_id]
        journal_lines = [ln for ln in lines_raw if ln["gl_journal_id"] == journal_id]
        journal_lines.sort(key=lambda ln: (ln.get("sort_order") or 0))

        for line in journal_lines:
            acct = account_map.get(line.get("account_id", ""), {})
            debit = _money(line.get("debit_amount"))
            credit = _money(line.get("credit_amount"))
            total_debits += debit
            total_credits += credit

            rows.append({
                "journal_id": journal["id"],
                "journal_date": journal["journal_date"],
                "journal_description": journal.get("description") or "",
                "source_type": journal.get("source_type") or "",
                "source_label": _source_label(journal.get("source_type")),
                "line_id": line["id"],
                "account_id": line.get("account_id") or "",
                "account_code": acct.get("code"),
                "account_name": acct.get("name") or "",
                "account_type": acct.get("type") or "",
                "line_description": line.get("description") or "",
                "debit_amount": _amount_out(debit),
                "credit_amount": _amount_out(credit),
            })

    return {
        "organisation_id": organisation_id,
        "date_from": date_from,
        "date_to": date_to,
        "rows": rows,
        "summary": {
            "journal_count": len(journal_ids),
            "line_count": len(rows),
            "total_debits": _amount_out(total_debits),
            "total_credits": _amount_out(total_credits),
        },
    }


# ── Exports ───────────────────────────────────────────────────────────────────

_CSV_HEADERS = [
    "Date",
    "Journal ID",
    "Source",
    "Journal Description",
    "Account Code",
    "Account Name",
    "Account Type",
    "Line Description",
    "Debit",
    "Credit",
]


def _row_to_csv(row: dict) -> list:
    return [
        row["journal_date"],
        row["journal_id"],
        row["source_label"],
        row["journal_description"],
        row["account_code"] or "",
        row["account_name"],
        row["account_type"],
        row["line_description"],
        f"{row['debit_amount']:.2f}" if row["debit_amount"] else "",
        f"{row['credit_amount']:.2f}" if row["credit_amount"] else "",
    ]


def transaction_report_csv(report: dict) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_CSV_HEADERS)
    for row in report["rows"]:
        writer.writerow(_row_to_csv(row))
    s = report["summary"]
    writer.writerow([])
    writer.writerow(["", "", "", "", "", "", "", "Totals", f"{s['total_debits']:.2f}", f"{s['total_credits']:.2f}"])
    return buf.getvalue().encode("utf-8-sig")


def transaction_report_text(report: dict) -> bytes:
    lines: list[str] = []
    lines.append(f"Transactions Report")
    lines.append(f"Period: {report['date_from']} to {report['date_to']}")
    lines.append("")
    lines.append(f"{'Date':<12} {'Source':<10} {'Account':<40} {'Description':<50} {'Debit':>14} {'Credit':>14}")
    lines.append("-" * 144)
    for row in report["rows"]:
        acct = f"{row['account_code'] or ''} {row['account_name']}".strip()
        desc = (row["line_description"] or row["journal_description"] or "")[:50]
        debit = f"{row['debit_amount']:>14.2f}" if row["debit_amount"] else " " * 14
        credit = f"{row['credit_amount']:>14.2f}" if row["credit_amount"] else " " * 14
        lines.append(f"{row['journal_date']:<12} {row['source_label']:<10} {acct:<40} {desc:<50} {debit} {credit}")
    s = report["summary"]
    lines.append("-" * 144)
    lines.append(f"{'':12} {'':10} {'':40} {'Totals':<50} {s['total_debits']:>14.2f} {s['total_credits']:>14.2f}")
    lines.append("")
    lines.append(f"Journals: {s['journal_count']}   Lines: {s['line_count']}")
    return "\n".join(lines).encode("utf-8")


def transaction_report_xlsx(report: dict) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel export") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1F497D")

    ws.append(_CSV_HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in report["rows"]:
        ws.append(_row_to_csv(row))

    s = report["summary"]
    ws.append([])
    total_row = ["", "", "", "", "", "", "", "Totals", s["total_debits"], s["total_credits"]]
    ws.append(total_row)
    last = ws.max_row
    for col in (9, 10):
        cell = ws.cell(row=last, column=col)
        cell.font = Font(bold=True)
        cell.number_format = "#,##0.00"

    col_widths = [12, 38, 10, 45, 14, 35, 12, 45, 14, 14]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 2):
        for cell in row[8:10]:
            cell.number_format = "#,##0.00"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
