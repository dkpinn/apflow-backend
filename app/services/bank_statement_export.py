from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Optional

ZERO = Decimal("0.00")
MONEY = Decimal("0.01")

_CSV_HEADERS = ["Date", "Description", "Reference", "Counterparty", "Debit", "Credit", "Balance", "Status"]


def _money(value: Any) -> Decimal:
    if value in (None, ""):
        return ZERO
    try:
        return Decimal(str(value)).quantize(MONEY, rounding=ROUND_HALF_UP)
    except Exception:
        return ZERO


def _validate_dates(date_from: str, date_to: str) -> None:
    try:
        start = date.fromisoformat(date_from)
        end = date.fromisoformat(date_to)
    except ValueError as exc:
        raise ValueError("Dates must use YYYY-MM-DD format") from exc
    if start > end:
        raise ValueError("From date must be on or before To date")


def generate_bank_statement_report(
    db,
    account_id: str,
    organisation_id: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
) -> dict:
    if date_from and date_to:
        _validate_dates(date_from, date_to)

    # Fetch bank account info
    acct_result = (
        db.table("bank_accounts")
        .select("id, name, account_number_mask, currency")
        .eq("id", account_id)
        .eq("organisation_id", organisation_id)
        .maybe_single()
        .execute()
    )
    account = acct_result.data or {}

    # Fetch lines
    query = (
        db.table("bank_statement_lines")
        .select(
            "line_date, description, reference, counterparty, "
            "debit_amount, credit_amount, balance_amount, match_status, posting_status"
        )
        .eq("bank_account_id", account_id)
        .eq("organisation_id", organisation_id)
        .order("line_date", desc=False)
        .order("id", desc=False)
    )
    if date_from:
        query = query.gte("line_date", date_from)
    if date_to:
        query = query.lte("line_date", date_to)

    result = query.execute()
    lines = result.data or []

    rows = []
    total_debits = ZERO
    total_credits = ZERO
    for line in lines:
        debit = _money(line.get("debit_amount"))
        credit = _money(line.get("credit_amount"))
        balance = _money(line.get("balance_amount"))
        total_debits += debit
        total_credits += credit

        match_status = line.get("match_status") or ""
        posting_status = line.get("posting_status") or ""
        status = posting_status if posting_status else match_status

        rows.append({
            "line_date": str(line.get("line_date") or ""),
            "description": line.get("description") or "",
            "reference": line.get("reference") or "",
            "counterparty": line.get("counterparty") or "",
            "debit_amount": float(debit) if debit else None,
            "credit_amount": float(credit) if credit else None,
            "balance_amount": float(balance) if balance else None,
            "status": status,
        })

    return {
        "account_name": account.get("name") or "",
        "account_number_mask": account.get("account_number_mask") or "",
        "currency": account.get("currency") or "ZAR",
        "date_from": date_from or "",
        "date_to": date_to or "",
        "rows": rows,
        "summary": {
            "total_debits": float(total_debits),
            "total_credits": float(total_credits),
            "row_count": len(rows),
        },
    }


def _row_to_list(row: dict) -> list:
    return [
        row["line_date"],
        row["description"],
        row["reference"],
        row["counterparty"],
        f"{row['debit_amount']:.2f}" if row["debit_amount"] is not None else "",
        f"{row['credit_amount']:.2f}" if row["credit_amount"] is not None else "",
        f"{row['balance_amount']:.2f}" if row["balance_amount"] is not None else "",
        row["status"],
    ]


def bank_statement_csv(report: dict) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_CSV_HEADERS)
    for row in report["rows"]:
        writer.writerow(_row_to_list(row))
    s = report["summary"]
    writer.writerow([])
    writer.writerow(["", "", "", "Totals", f"{s['total_debits']:.2f}", f"{s['total_credits']:.2f}", "", ""])
    return buf.getvalue().encode("utf-8-sig")


def bank_statement_xlsx(report: dict) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel export") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bank Statement"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1F497D")

    ws.append(_CSV_HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in report["rows"]:
        ws.append([
            row["line_date"],
            row["description"],
            row["reference"],
            row["counterparty"],
            row["debit_amount"],
            row["credit_amount"],
            row["balance_amount"],
            row["status"],
        ])

    s = report["summary"]
    ws.append([])
    total_row = ["", "", "", "Totals", s["total_debits"], s["total_credits"], "", ""]
    ws.append(total_row)
    last = ws.max_row
    for col in (5, 6, 7):
        cell = ws.cell(row=last, column=col)
        cell.font = Font(bold=True)
        cell.number_format = "#,##0.00"

    col_widths = [12, 45, 20, 30, 14, 14, 14, 15]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 2):
        for cell in (row[4], row[5], row[6]):  # Debit, Credit, Balance columns
            cell.number_format = "#,##0.00"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
