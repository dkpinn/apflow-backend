import io
from decimal import Decimal

import pytest

from app.routers.reports import _ensure_reports_view
from app.services.vat_report import (
    allocate_amount_by_weights,
    allocate_invoice_vat,
    generate_vat_report,
    vat_report_csv,
    vat_report_text,
    vat_report_xlsx,
)


class _Response:
    def __init__(self, data=None):
        self.data = data or []


class _Query:
    def __init__(self, rows):
        self.rows = list(rows or [])
        self.filters = []
        self.in_filters = []
        self.lte_filters = []
        self._limit = None
        self.orders = []

    def select(self, *_args, **_kwargs):
        return self

    def eq(self, field, value):
        self.filters.append((field, value))
        return self

    def in_(self, field, values):
        self.in_filters.append((field, {str(value) for value in values}))
        return self

    def lte(self, field, value):
        self.lte_filters.append((field, value))
        return self

    def limit(self, value):
        self._limit = value
        return self

    def order(self, field, *_args, **_kwargs):
        self.orders.append(field)
        return self

    def execute(self):
        rows = self.rows
        for field, value in self.filters:
            rows = [row for row in rows if row.get(field) == value]
        for field, values in self.in_filters:
            rows = [row for row in rows if str(row.get(field)) in values]
        for field, value in self.lte_filters:
            rows = [row for row in rows if row.get(field) <= value]
        for field in reversed(self.orders):
            rows = sorted(rows, key=lambda row: (row.get(field) is None, row.get(field)))
        if self._limit is not None:
            rows = rows[: self._limit]
        return _Response(rows)


class _DB:
    def __init__(self, tables):
        self.tables = tables

    def table(self, name):
        return _Query(self.tables.get(name, []))


def _tables():
    return {
        "accounts": [
            {
                "id": "vat-control",
                "organisation_id": "org-1",
                "system_key": "vat_control",
                "code": "8100",
                "name": "VAT Control",
            }
        ],
        "gl_journals": [
            {
                "id": "opening",
                "organisation_id": "org-1",
                "status": "posted",
                "source_type": "manual",
                "source_id": None,
                "journal_date": "2025-12-31",
                "description": "Opening VAT",
                "created_at": "2025-12-31T10:00:00Z",
            },
            {
                "id": "purchase",
                "organisation_id": "org-1",
                "status": "posted",
                "source_type": "invoice",
                "source_id": "invoice-1",
                "journal_date": "2026-01-10",
                "description": "Supplier invoice",
                "created_at": "2026-01-10T10:00:00Z",
            },
            {
                "id": "sale",
                "organisation_id": "org-1",
                "status": "posted",
                "source_type": "manual",
                "source_id": None,
                "journal_date": "2026-01-20",
                "description": "Output VAT",
                "created_at": "2026-01-20T10:00:00Z",
            },
            {
                "id": "draft",
                "organisation_id": "org-1",
                "status": "draft",
                "source_type": "manual",
                "source_id": None,
                "journal_date": "2026-01-25",
                "description": "Draft VAT",
                "created_at": "2026-01-25T10:00:00Z",
            },
        ],
        "gl_journal_lines": [
            {
                "id": "opening-line",
                "organisation_id": "org-1",
                "gl_journal_id": "opening",
                "account_id": "vat-control",
                "description": "Opening VAT",
                "debit_amount": 0,
                "credit_amount": 50,
                "created_at": "2025-12-31T10:00:00Z",
                "sort_order": 0,
            },
            {
                "id": "purchase-line",
                "organisation_id": "org-1",
                "gl_journal_id": "purchase",
                "account_id": "vat-control",
                "description": "Example Supplier - INV-1 - VAT",
                "debit_amount": 15,
                "credit_amount": 0,
                "created_at": "2026-01-10T10:00:00Z",
                "sort_order": 1,
            },
            {
                "id": "sale-line",
                "organisation_id": "org-1",
                "gl_journal_id": "sale",
                "account_id": "vat-control",
                "description": "Output VAT adjustment",
                "debit_amount": 0,
                "credit_amount": 40,
                "created_at": "2026-01-20T10:00:00Z",
                "sort_order": 0,
            },
            {
                "id": "draft-line",
                "organisation_id": "org-1",
                "gl_journal_id": "draft",
                "account_id": "vat-control",
                "description": "Draft",
                "debit_amount": 0,
                "credit_amount": 999,
                "created_at": "2026-01-25T10:00:00Z",
                "sort_order": 0,
            },
        ],
        "invoices_extracted": [
            {
                "id": "invoice-1",
                "organisation_id": "org-1",
                "supplier_id": "supplier-1",
                "supplier_name_extracted": "Example Supplier OCR",
                "vat_number_extracted": "4000000001",
                "invoice_number": "INV-1",
                "subtotal": 100,
                "tax_amount": 15,
                "total_amount": 115,
            }
        ],
        "suppliers": [
            {
                "id": "supplier-1",
                "organisation_id": "org-1",
                "supplier_name": "Example Supplier",
                "trading_name": None,
                "vat_number": "4000000001",
            }
        ],
        "invoice_line_items": [
            {
                "id": "item-full",
                "organisation_id": "org-1",
                "invoice_extracted_id": "invoice-1",
                "line_total": 70,
                "tax_amount": None,
                "vat_treatment": "full",
            },
            {
                "id": "item-blocked",
                "organisation_id": "org-1",
                "invoice_extracted_id": "invoice-1",
                "line_total": 30,
                "tax_amount": None,
                "vat_treatment": "blocked",
            },
        ],
        "organisation_users": [],
    }


def test_vat_report_reconciles_opening_running_total_and_claimability():
    report = generate_vat_report(
        _DB(_tables()),
        organisation_id="org-1",
        date_from="2026-01-01",
        date_to="2026-01-31",
    )

    assert report["summary"] == {
        "opening_balance": 50.0,
        "output_vat": 40.0,
        "posted_input_vat": 15.0,
        "allowable_input_vat": 10.5,
        "blocked_input_vat": 4.5,
        "period_vat_payable_refundable": 29.5,
        "period_gl_movement": 25.0,
        "closing_vat_control_balance": 75.0,
        "calculated_vat_position": 79.5,
        "historical_claimability_variance": 4.5,
    }
    assert len(report["rows"]) == 2
    assert report["rows"][0]["supplier"] == "Example Supplier"
    assert report["rows"][0]["supplier_vat_number"] == "4000000001"
    assert report["rows"][0]["invoice_number"] == "INV-1"
    assert report["rows"][0]["gross_amount"] == 115.0
    assert report["rows"][0]["running_total"] == 35.0
    assert report["rows"][1]["running_total"] == 75.0
    assert report["warnings"][0]["code"] == "historical_claimability_variance"


def test_vat_allocation_handles_mixed_treatments_and_non_vat_supplier():
    lines = [
        {"id": "full", "line_total": 70, "tax_amount": None, "vat_treatment": "full"},
        {"id": "blocked", "line_total": 30, "tax_amount": None, "vat_treatment": "blocked"},
        {"id": "zero", "line_total": 10, "tax_amount": None, "vat_treatment": "zero_rated"},
    ]

    registered = allocate_invoice_vat(
        invoice_tax=15,
        line_items=lines,
        supplier_has_vat_number=True,
    )
    assert registered["claimable_tax"] == 10.5
    assert registered["blocked_tax"] == 4.5

    unregistered = allocate_invoice_vat(
        invoice_tax=15,
        line_items=lines,
        supplier_has_vat_number=False,
    )
    assert unregistered["claimable_tax"] == 0
    assert unregistered["blocked_tax"] == 15


def test_weighted_allocation_preserves_rounding_total():
    shares = allocate_amount_by_weights(1, [1, 1, 1])
    assert shares == [Decimal("0.33"), Decimal("0.33"), Decimal("0.34")]
    assert sum(shares) == Decimal("1.00")


def test_vat_report_exports_match_detail_rows():
    report = generate_vat_report(
        _DB(_tables()),
        organisation_id="org-1",
        date_from="2026-01-01",
        date_to="2026-01-31",
    )

    csv_text = vat_report_csv(report).decode("utf-8-sig")
    assert "Supplier VAT Number" in csv_text
    assert "Example Supplier" in csv_text

    text = vat_report_text(report).decode("utf-8")
    assert "Calculated VAT position\t79.50" in text
    assert "INV-1" in text

    workbook_bytes = vat_report_xlsx(report)
    assert workbook_bytes[:2] == b"PK"
    assert len(io.BytesIO(workbook_bytes).getvalue()) > 1000
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(workbook_bytes))
    assert workbook["Summary"]["B2"].number_format == "yyyy-mm-dd"
    assert workbook["VAT Detail"]["A2"].number_format == "yyyy-mm-dd"


def test_blocked_invoice_without_vat_control_line_still_appears_in_summary():
    tables = _tables()
    tables["gl_journals"].append({
        "id": "blocked-purchase",
        "organisation_id": "org-1",
        "status": "posted",
        "source_type": "invoice",
        "source_id": "invoice-blocked",
        "journal_date": "2026-01-15",
        "description": "Blocked supplier invoice",
        "created_at": "2026-01-15T10:00:00Z",
    })
    tables["invoices_extracted"].append({
        "id": "invoice-blocked",
        "organisation_id": "org-1",
        "supplier_id": "supplier-1",
        "supplier_name_extracted": "Example Supplier OCR",
        "vat_number_extracted": "4000000001",
        "invoice_number": "INV-BLOCKED",
        "subtotal": 100,
        "tax_amount": 15,
        "total_amount": 115,
    })
    tables["invoice_line_items"].append({
        "id": "item-all-blocked",
        "organisation_id": "org-1",
        "invoice_extracted_id": "invoice-blocked",
        "line_total": 100,
        "tax_amount": None,
        "vat_treatment": "blocked",
    })

    report = generate_vat_report(
        _DB(tables),
        organisation_id="org-1",
        date_from="2026-01-01",
        date_to="2026-01-31",
    )

    assert report["summary"]["allowable_input_vat"] == 10.5
    assert report["summary"]["blocked_input_vat"] == 19.5
    assert len(report["rows"]) == 2


def test_reports_permission_allows_roles_or_explicit_permission():
    tables = _tables()
    tables["organisation_users"] = [
        {"organisation_id": "org-1", "user_id": "owner", "status": "active", "role": "owner", "permissions": {}},
        {
            "organisation_id": "org-1",
            "user_id": "viewer",
            "status": "active",
            "role": "viewer",
            "permissions": {"reports_view": True},
        },
        {"organisation_id": "org-1", "user_id": "blocked", "status": "active", "role": "viewer", "permissions": {}},
    ]
    db = _DB(tables)

    _ensure_reports_view(db, "owner", "org-1")
    _ensure_reports_view(db, "viewer", "org-1")
    with pytest.raises(Exception) as exc:
        _ensure_reports_view(db, "blocked", "org-1")
    assert getattr(exc.value, "status_code", None) == 403


def test_vat_report_rejects_invalid_date_range():
    with pytest.raises(ValueError, match="From date"):
        generate_vat_report(
            _DB(_tables()),
            organisation_id="org-1",
            date_from="2026-02-01",
            date_to="2026-01-31",
        )
